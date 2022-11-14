from tkinter import *
from openpyxl import load_workbook
from tkinter import filedialog as fd
from tkinter import messagebox


root = Tk()
root.title('Data Generator')
root.config(pady=10)

e1 = Entry(root, width=30)
e1.grid(row=1, column=2, pady=5, padx=5)

e2 = Entry(root, width=30)
e2.grid(row=2, column=2, pady=5, padx=5)

e3 = Entry(root, width=30)
e3.grid(row=3, column=2, pady=5, padx=5)


l1 = Label(root, text='Copy to')
l1.grid(row=1, column=1, pady=5, padx=5, sticky='W')

l2 = Label(root, text='Copy from')
l2.grid(row=2, column=1, pady=5, padx=5, sticky='W')

l3 = Label(root, text='Save to')
l3.grid(row=3, column=1, pady=5, padx=5, sticky='W')


def select_file1():
    file_name = fd.askopenfilename(title='Select a file')
    e1.insert(0, file_name)


def select_file2():
    file_name = fd.askopenfilename(title='Select a file')
    e2.insert(0, file_name)


def select_file3():
    file_name = fd.askdirectory(title='Select a folder')
    e3.insert(0, file_name)


b1 = Button(root, text='Browse', command=select_file1)
b1.grid(row=1, column=3, pady=5, padx=5, sticky='W')

b2 = Button(root, text='Browse', command=select_file2)
b2.grid(row=2, column=3, pady=5, padx=5, sticky='W')

b3 = Button(root, text='Browse', command=select_file3)
b3.grid(row=3, column=3, pady=5, padx=5, sticky='W')


def copy_data():
    wb1 = load_workbook(e1.get())
    wb2 = load_workbook(e2.get())
    ws1 = wb1['Sheet name']
    data = []
    data_needed = wb2.sheetnames
    # This for loop is to extract data from the sheets with values and save it in a list
    for sheet in data_needed:
        ws2 = wb2[sheet]
        for row in ws2.iter_rows(min_row=num1, min_col=num2, values_only=True):
            row = list(row)
            if  row[12] == None :
                continue
            elif isinstance(row[12],str):
                continue
            elif row[12] >5:
                row.insert(0, sheet[-5:])
                data.append(row)

    file_to_be_added = ws1.max_row
    for r, t in enumerate(data, start=file_to_be_added+1):
        for c, d in enumerate(t, start=2):
            ws1.cell(row=r, column=c, value=d)
    wb1.save(e3.get() + '/New Document.xlsx')

    messagebox.showinfo("Data copied", "A new file has been generated")

    e1.delete(0, END)
    e2.delete(0, END)
    e3.delete(0, END)

b4 = Button(root, text='Run', width=10, command=copy_data)
b4.grid(row=4, column=2, pady=5)

root.mainloop()
